from pathlib import Path
import ast
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image,ImageDraw,ImageFont
R=Path.cwd();tree=ast.parse((R/'scripts/create_interview_presentation.py').read_text(encoding='utf-8-sig'));base=[];examples=[]
for n in tree.body:
 if isinstance(n,ast.Assign) and any(isinstance(t,ast.Name) and t.id=='slides' for t in n.targets):base=ast.literal_eval(n.value)
 if isinstance(n,ast.Expr) and isinstance(n.value,ast.Call) and isinstance(n.value.func,ast.Attribute) and isinstance(n.value.func.value,ast.Name) and n.value.func.value.id=='slides' and n.value.func.attr=='extend':examples=ast.literal_eval(n.value.args[0])
slides=[base[i] for i in [0,1,2,3,4]]
slides += [('System architecture','Local components exchange traceable text blocks and explicit processing states.','ARCH','Describe the arrows in order. Browser requests go through session and file validation. The runtime orchestrates real adapters, translators and release checks. Raw evidence stays local; the downstream LLM is a future integration after checked export.')]
slides += [base[5],base[6]]
slides += [('Reading uploaded files','Native text and OCR converge into the same block contract.','PDF: read native characters; render scanned pages with PDFium for Tesseract OCR.\nPPTX: inspect OOXML slide shapes, tables, notes and relevant embedded media.\nDOCX: inspect OOXML paragraphs, tables, headers / footers and embedded media.\nSVG / EMF: convert supported images locally before recognition.\nOutput: text + document / block ID + source location + quality warnings.', 'OOXML is the ZIP/XML representation of Word and PowerPoint documents. Extraction does not require copying text from the visible Office window. OCR recognizes characters from pixels; it does not classify PII. Unsupported conversions remain explicit review issues.'),base[7],base[8],base[9],base[10]]
slides += [('PII detection: from text to findings','Different signals address formatted identifiers and natural-language names.','1. Read each normalized block while preserving Unicode character positions.\n2. Recognize likely people using the local spaCy English NER model.\n3. Apply patterns, nearby labels, date checks and relevant checksum validators.\n4. Resolve document-local names / aliases and retain repeated occurrences.\n5. Return category, block reference and exact start / end offsets.', 'Local NER is not an external LLM. Email patterns, context such as DOB, and validators have different strengths. The detector attempts to reject ordinary workflow labels. False positives and OCR-induced misses remain possible; independent evaluation is necessary.'),base[12],base[13],base[14]]
slides += [('Backend processing and response','The application receives one consistent result from the runtime.','Upload bytes → file validation → per-request extraction and detection context.\nAdapters reconcile confidence scales, document IDs and block references.\nRedaction produces a masked preview; verification produces the release state.\nBrowser polls the job for stage progress, counts, locations and warnings.\nDownload route enforces approval; original text requires explicit local inspection.', 'Code references: web.py, runtime.py, ingestion/orchestrator.py, redaction/redactor.py and redaction/gate.py. The server uses local workers and in-memory jobs. Original values and reversible mappings must not be treated as public audit output.'),base[15],base[16],base[17],base[18],base[20]]
slides += examples[1:5]
slides += [('Implementation novelty','Our contribution is a traceable release decision across mixed evidence formats.','A shared evidence contract joins native text, OCR, source locations and quality.\nContextual rules + local NER + validators address different PII failure modes.\nOverlap-safe masking and original-span checks reduce partial-value leakage.\nThe release gate requires both masked content and processing completeness.\nJudges can trace a finding to its source and see why export was allowed or blocked.', 'Do not claim that OCR, regex or NER is a newly invented algorithm. Differentiation is the implemented combination, traceability and fail-closed behavior. Peers receive the same problem; compare demonstrable behavior rather than claiming unverified superiority. Our bounded synthetic results support specific checks, not universal accuracy.')]
prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
navy='12284C';gold='C8A34B';ink='3A4557';muted='616F87'
def rect(sl,x,y,w,h,fill,line=None):
 sh=sl.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(fill);sh.line.color.rgb=RGBColor.from_string(line or fill);return sh
def text(sl,x,y,w,h,t,size=20,color=ink,font='Calibri',bold=False):
 sh=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));sh.text_frame.word_wrap=True
 for i,line in enumerate(t.split('\n')):
  p=sh.text_frame.paragraphs[0] if i==0 else sh.text_frame.add_paragraph();p.text=line;p.font.name=font;p.font.size=Pt(size);p.font.color.rgb=RGBColor.from_string(color);p.font.bold=bold;p.space_after=Pt(3 if size<20 else 18)
 return sh
for i,(title,claim,body,notes) in enumerate(slides):
 title=title.replace('Cadence:','Caddet:').replace('Role: ',''); title={'The problem we were given':'Problem statement','What the judges asked to see':'Case study requirements','The complete journey':'End-to-end pipeline','How the code fits together':'Architecture: component responsibilities','What we classify':'PII categories','What the tests currently establish':'Validation results','The cases we deliberately challenge':'Test coverage','What changed during integration':'Integration improvements'}.get(title,title);title=title[:1].upper()+title[1:]
 sl=prs.slides.add_slide(prs.slide_layouts[6]);rect(sl,0,0,13.333,7.5,'FFFFFF');rect(sl,0,0,13.333,1.05,navy);text(sl,.55,.1,12.1,.55,title,26,'FFFFFF','Georgia');text(sl,.57,.65,12,.3,'CADDET | TEXT ANALYSIS AND PII DETECTION',10,gold);text(sl,.6,1.35,12.1,.8,claim,21,navy,'Georgia')
 if body=='ARCH':
  nodes=[(.65,2.6,'Browser interface','Upload / poll / inspect\nDownload checked text'),(4.65,2.6,'Local server + runtime','Session / file validation\nStage orchestration'),(8.65,2.6,'Extraction + OCR','PDF / PPTX / DOCX\nText, locations, warnings'),(8.65,4.65,'PII detection','Local NER + rules\nTyped character spans'),(4.65,4.65,'Redaction','Validate / union spans\nStable typed tokens'),(.65,4.65,'Verification + gate','Rescan / completeness\nApprove or block')]
  for x,y,a,b in nodes:rect(sl,x,y,3.6,1.35,'F2F4F7','D7DEE7');text(sl,x+.15,y+.12,3.3,.35,a,19,navy,'Georgia');text(sl,x+.15,y+.6,3.3,.7,b,16,muted)
  for x in [4.25,8.25]:text(sl,x,2.95,.4,.6,'→',24,navy);text(sl,x,4.95,.4,.6,'←',24,navy)
  text(sl,10.25,4,.5,.55,'↓',24,navy);text(sl,.8,6.3,11.8,.55,'Approved → checked text export     |     Uncertain / incomplete → local review',18,navy)
 else:
  rect(sl,.6,2.4,.045,4.0,gold);text(sl,.9,2.4,11.65,4.25,body,22)
 rect(sl,0,7.08,13.333,.42,navy);text(sl,.55,7.1,11,.26,'Caddet | Team 57',10,'B0B7C3');text(sl,12.3,7.1,.5,.26,str(i+1),10,gold)
 sl.notes_slide.notes_text_frame.text=notes
out=R/'presentation/Caddet_Formal_Presentation.pptx';prs.save(out)
preview=R/'tmp/formal-preview';preview.mkdir(exist_ok=True)
for i,sl in enumerate(prs.slides):
 im=Image.new('RGB',(1280,720),'white');d=ImageDraw.Draw(im)
 for sh in sl.shapes:
  x,y,w,h=[round(v/914400*96) for v in (sh.left,sh.top,sh.width,sh.height)]
  if sh.shape_type==1:d.rectangle((x,y,x+w,y+h),fill=tuple(sh.fill.fore_color.rgb),outline=tuple(sh.line.color.rgb))
  if not sh.has_text_frame or not sh.text:continue
  yy=y+4
  for p in sh.text_frame.paragraphs:
   fname='georgia.ttf' if p.font.name=='Georgia' else ('calibrib.ttf' if p.font.bold else 'calibri.ttf');f=ImageFont.truetype('C:/Windows/Fonts/'+fname,round(p.font.size.pt*96/72));lines=[];line=''
   for word in p.text.split():
    candidate=(line+' '+word).strip()
    if line and d.textlength(candidate,font=f)>w-20:lines.append(line);line=word
    else:line=candidate
   lines.append(line)
   for line in lines:d.text((x+9,yy),line,font=f,fill=tuple(p.font.color.rgb));yy+=round(p.font.size.pt*96/72*1.16)
   yy+=round((3 if p.font.size.pt<20 else 18)*96/72)
  assert yy<=y+h+25,(i+1,sh.text[:50],yy,y+h)
 im.save(preview/f'{i+1:02}.png')
for start in range(0,len(slides),6):
 im=Image.new('RGB',(960,810),'#ddd')
 for j in range(start,min(start+6,len(slides))):
  pic=Image.open(preview/f'{j+1:02}.png');pic.thumbnail((480,270));im.paste(pic,((j-start)%2*480,(j-start)//2*270))
 im.save(preview/f'contact-{start}.jpg')
print('Created',len(slides),'slides:',out)
