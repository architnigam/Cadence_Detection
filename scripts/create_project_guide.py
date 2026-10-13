"""Create the beginner Word guide and editable case-study presentation."""
import bootstrap
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.util import Inches as PInches,Pt as PPt
from pptx.dml.color import RGBColor as PColor
from pptx.enum.shapes import MSO_SHAPE

ROOT=bootstrap.ROOT
testpath=ROOT/'output/reports/final_tests.json'
tests=json.loads(testpath.read_text(encoding='utf-8')) if testpath.exists() else {}
pages=[]
def page(title,sections):pages.append((title,sections))

page('Cadence Detection | Team learning guide',[
('What the project does','Cadence uses business documents with an analysis agent. The case study says personally identifiable information (PII) must be removed before document text reaches an LLM. Our application reads PDF, PowerPoint and Word files locally, identifies sensitive text, replaces it with labels, and verifies whether text export can be enabled.'),
('A simple example','Invented input: Name: Noor Example; email: noor@example.invalid. Checked output: Name: [PERSON_1]; email: [EMAIL_1]. The business sentence remains readable while the identity is removed. A document with unreadable content is held for review rather than declared safe.'),
('What to learn first','Read the shared vocabulary, then your role chapter, then the run and demonstration instructions. Everyone should understand the complete flow even if they work on only one component. The five roles below describe ownership by component; the final integration work joins the five branches and fixes their mismatches.'),
('What is delivered','A real local upload application, command-line pipeline, local OCR and name recognition, a regression suite, evaluation/report tooling and an editable stakeholder deck. The pipeline produces sanitized plain text, not redacted copies of the original PDF/PPTX/DOCX. Original photographs, signatures and image pixels are not exported.'),
('Read results honestly','A passing test proves the behavior tested by that fixture. It does not prove perfect recognition on every future document. The supplied synthetic text benchmark has 18 labeled occurrences. Full source PII recall and the case-study structure target still need independently reviewed source annotations. OCR warnings are intentionally visible and block release.'),
('Prepared','14 October 2026. Component ownership follows the team allocation: Person 1 ingestion; Person 2 extraction; Person 3 detection; Person 4 redaction and UI; Person 5 evaluation. Replace these labels with your team members\' names during presentation preparation.')])
page('1 | Shared vocabulary for everyone',[
('PII and LLM','PII is information that identifies a person directly or in combination: names, contact details, employee identifiers, birth dates, addresses and personal financial or national identifiers. An LLM is a language model that processes text. This project stops unverified text before that boundary; it does not need to call an external LLM to demonstrate the control.'),
('Extraction, OCR and NER','Extraction means recovering readable text and its location from a file. Native text is already stored as text. OCR recognizes letters from pixels, such as a scan or screenshot. NER means named-entity recognition: a local language model proposes person-name spans in prose. NER makes mistakes; it is combined with rules and verification.'),
('Block, finding, offset and confidence','A block is a paragraph, table cell, slide text shape or OCR line. A finding says which kind of PII occurs in a block and which characters it occupies. A start offset includes its character; the end offset excludes its character. Python counts Unicode code points, so Zoë and emoji must not be measured as UTF-8 bytes. Confidence is an estimate used for review, not a calibrated probability.'),
('Masking, quarantine and release','Masking replaces a sensitive span with a typed token. Stable tokens reuse the same identity label within one document. Quarantine means the document needs review and export is blocked. A completed pipeline with safe_for_llm means its configured checks passed; it is not a claim of universal PII certainty.'),
('Git, branches and tests','Git records source history. A branch lets someone develop a component separately. A merge combines changes; it does not automatically connect function signatures or data models. Unit tests check small components. Integration tests run cooperating components. Acceptance tests exercise the scenario from the case-study problem statement.'),
('API, adapter and dependency','An API is the interface one component exposes to another. An adapter translates between APIs. A dependency is an external library or model. Python packages, the OCR language file and the English NER model must be installed before the full local demonstration. Runtime processing never downloads models or uploads evidence.')])
page('2 | Case study and the complete flow',[
('Required stakeholder outcomes','The assignment requests a PowerPoint walkthrough lasting 10-15 minutes: a functional solution blueprint, detailed system architecture, extraction/detection logic, reliable outcomes and a live demonstration. It sets at least 80% structure/context retention and PII coverage approaching 100% as goals. Approximately three weeks from release is a planning duration, not a known calendar deadline.'),
('Evidence supplied','The organizational pack has 9 slides with a directory, organization charts and RACI tables. The risk policy has 35 scanned pages, including narrative incidents, rosters, screenshots and an identity badge. The training DOCX has 16 tables and 41 media entries, including PNG, SVG and EMF content. Different documents use different role-holder examples; they must not be treated as one authoritative live directory.'),
('The processing sequence','1. Upload a supported file locally. 2. Validate size, signature, format and integrity. 3. Extract native text and OCR images. 4. Keep blocks, source locations, reading order and confidence. 5. Detect and classify PII. 6. Union overlapping sensitive ranges and mask them. 7. Rescan masked text and check original detected spans. 8. Enable checked text export or retain the document for local review.'),
('The local boundary','The application listens on loopback. It has no external image, font or analytics requests. Upload requests require a per-launch session token and valid local origin/host. Two workers bound processing concurrency. Completed review jobs expire after 20 minutes when a new upload arrives; restart clears the in-memory session. This is a local demonstration service, not an internet deployment.'),
('How the team connects','Person 1 controls the lifecycle and release state. Person 2 supplies text and completeness. Person 3 supplies spans. Person 4 applies spans and verifies output. Person 5 independently evaluates the claims. The runtime module supplies real adapters for the existing orchestration contracts; the older stubs remain only for tests that verify fail-closed behavior.')])

roles=[
('3 | Person 1: architecture, ingestion and orchestration',[
('Your responsibility','Make sure every file enters the same controlled process. Draw the functional and technical diagrams, define component contracts, validate inputs and manage success/failure states. An upload must never bypass a failed step just because another component returned some text.'),
('Code you should know','src/cadence_detection/ingestion/validator.py checks supported formats, signatures, size, corruption and encryption. contracts.py defines manifests, normalized blocks, findings and responses. orchestrator.py calls extraction, detection, redaction and rescan adapters. runtime.py constructs the real adapters used by the CLI and UI.'),
('Input and output','Input is a file path or bytes plus a display filename. A manifest contains the full SHA-256 content hash, format and size. Output contains completed/quarantined status, a release decision, counts and audit information. A quarantined public response has no releasable sanitized text.'),
('Important integration work','The extraction component originally used short hashes and 0-100 OCR confidence. Orchestration uses full hashes and 0-1 confidence. The adapter remaps every block ID consistently, divides confidence by 100 and preserves source geometry internally. Missing NER, extraction warnings and rescan failures block release. Review analysis may continue after incomplete extraction so users can inspect detections, but export remains blocked.'),
('Tests and demonstration','Explain the 12 file acceptance cases: three valid formats, three empty files, three corrupt files, unsupported text, a spoofed PDF and an encrypted PDF. Show the real pipeline path, then show a failure case that never calls the local downstream spy. Explain that a stub is a deliberately incomplete test double, not production processing.')]),
('4 | Person 2: document extraction and OCR',[
('Your responsibility','Recover text from every supported document without silently dropping pages, slides, tables or embedded evidence. Preserve where the text came from so reviewers can return to the source. Do not classify an unreadable image as an empty safe document.'),
('Code you should know','extraction/pipeline.py validates and routes extraction. pdf.py uses pypdf and PDFium for native text and rendered scans. office.py reads Word/PowerPoint OOXML, including tables, text boxes, notes, templates and image relationships. images.py converts supported media and maps boxes. ocr.py runs a local Tesseract backend; the JavaScript worker reuses its engine for a document.'),
('Native text versus scans','Native extraction is more direct than OCR and normally retains exact characters. A scanned PDF stores page pixels, so every policy page needs OCR. Word/PPTX can mix native text with screenshots. Small screenshots can receive a supplementary crop/upscale pass. The policy crop configuration applies to that policy, not arbitrary PDFs.'),
('Locations and structure','PDF OCR uses displayed-page points; native PDF rectangles retain unrotated coordinates and rotation metadata. Slide geometry uses EMUs. Word locations use XML paragraph/table paths because Word page numbers depend on rendering. OCR line/word geometry is retained, but semantic reconstruction of every scanned table cell is not claimed. Supplementary crops can repeat text; preserve location when reviewing duplicates.'),
('Warnings and tests','A low-confidence line, unreadable nonblank image, unsupported object or incomplete page makes extraction incomplete. Images are never passed to an external model. Test native/scanned/mixed PDF, rotation, nested tables, groups, headers/footers, notes, Unicode, SVG/EMF, limits, password/corruption and missing OCR. Show coverage for all 35 policy pages and 9 slides, and explain why coverage is not OCR accuracy.')]),
('5 | Person 3: PII detection and classification',[
('Your responsibility','Identify sensitive spans on the exact extracted text and classify them. Preserve offsets and source identity. Improve recall without treating every business number, date or organization name as personal data.'),
('Code you should know','detection/comprehensive.py is the document-level detector used by the integrated runtime. It combines local spaCy English NER, contextual rules, identifier patterns, a document-local name/alias directory and validation helpers. pii_detector.py remains the earlier baseline API with its historical regression tests; it is not the complete runtime detector.'),
('Recognized text categories','PERSON, EMAIL, PHONE, EMPLOYEE_ID, DIRECTOR_ID, VENDOR_PERSON_ID, DATE_OF_BIRTH, HOME_ADDRESS, NATIONAL_ID, TAX_ID, PASSPORT, BANK_ACCOUNT, PAYMENT_CARD and PARTIAL_IDENTIFIER. Names and dates use context; aliases cover initials in RACI tables. Card candidates use Luhn validation. Synthetic examples may intentionally violate real identifier issuance rules, so contextual recognition must not silently discard all teaching examples.'),
('Why document-level detection matters','A full name in a directory helps recognize its abbreviated occurrence later. A bare phone number in a directory cell must not require a nearby Phone label in the same cell. Repeated occurrences count separately. The name directory is confined to one document; it is not shared as a global identity database.'),
('Limits and tests','Confidence values are heuristic. Photograph/signature regions are not text recognizers; images are excluded from exports and uncertain visual evidence needs manual review. Test accented names, initials, hyphens, Unicode offsets, aliases, valid/invalid dates, financial checksums, repeated PII and negative business context. Missing the local NER model blocks the integrated release path instead of claiming full detection coverage.')]),
('6 | Person 4: redaction, privacy gate and application',[
('Your responsibility','Mask every detected sensitive range, retain readable surrounding business context, verify the result and provide a working local upload interface. A privacy gate must check evidence, not only display an approval banner.'),
('Code you should know','redaction/redactor.py validates offsets and unions overlapping spans. redaction/gate.py performs comprehensive rescanning, additional pattern checks, quality checks and original-span verification. web.py provides the local HTTP server, upload jobs, session protection and export endpoint. app/index.html is the actual frontend; app/app.py is a compatibility entry point.'),
('How masking works','Offsets refer to the original block. Replacements run from right to left so earlier offsets stay valid. When ranges 0-5 and 3-10 overlap, the whole 0-10 range is masked. Invalid offsets or unknown block/document IDs raise errors rather than being silently skipped and counted. Identical values reuse document-local typed tokens.'),
('Verification and review','The gate checks extraction completeness, model availability, OCR confidence, residual PII and whether known original spans survived. Its second pass uses the same local NER plus independent deterministic patterns; it is not a separately trained independent model. PDF creation timestamps are business metadata, not phone numbers. Raw matched values are excluded from gate report payloads.'),
('The user experience','Users upload PDF/PPTX/DOCX, follow actual stage progress, see category counts and masked previews, and inspect warning locations. Original text requires an explicit restricted local inspection action. Only completed approved jobs can download text; quarantined previews are not safe exports. There are no hardcoded leak counts or invented model upgrade claims. Test real upload/poll/download, unauthorized origins, failed jobs and blocked export.')]),
('7 | Person 5: evaluation, QA and submission',[
('Your responsibility','Check that project claims are supported by independent evidence. Build labels, run acceptance checks, report failures and prepare a clear stakeholder story. Do not use detector predictions as their own ground truth.'),
('Code you should know','evaluation/engine.py calculates exact occurrence/type metrics and diagnostic overlap metrics. pipeline_checks.py audits real release states. scripts/run_pipeline_checks.py calls runtime.run_for_qa for the 12 generated file cases. scripts/run_all_tests.py executes all unit/integration suites plus the earlier function-style redaction tests. Existing annotation and test-plan documents remain useful review aids.'),
('Precision, recall and F1','Precision = TP / (TP + FP): of detections made, how many were correct? Recall = TP / (TP + FN): of labeled sensitive occurrences, how many were found? F1 combines precision and recall. A wrong span boundary is a failure under exact matching even if overlap matching recognizes part of the value. Repeated mentions are separate occurrences.'),
('Structure and independent labels','The proposed structure rubric checks headings, reading order, block boundaries, table relationships and neighboring context, with N/A dimensions renormalized. The screenshot does not define this formula; explain it as a team measurement proposal. A page-processed count does not prove the 80% target. Complete and adjudicate source labels before reporting full source accuracy; no evaluation report can replace that work.'),
('Current evidence and presentation','The supplied synthetic text fixture contains 18 labeled occurrences; the integrated detector achieves 18 exact matches with zero false positives on it. This small bounded result must not be described as universal 100% accuracy. The final deck includes measured test evidence and limitations. Each owner presents their component and Person 5 coordinates the 10-15-minute rehearsal and final checks.')])]
for title,sections in roles:page(title,sections)
page('8 | Setup, run and demonstrate',[
('Fast run on this workstation','Use scripts/run_local.ps1 to select the bundled Python: powershell -NoProfile -ExecutionPolicy Bypass -File scripts/run_local.ps1 app. Open the printed localhost address. Choose one of the three actual source files and click Analyze document. OCR can take minutes; wait for the stage/status to finish.'),
('Standard Python setup','On another machine install Python 3.10+ and run python -m pip install -e ".[nlp,test]". Install the en_core_web_sm English model explicitly using the official spaCy model wheel or python -m spacy download en_core_web_sm when network access is authorized. Provision local Tesseract plus English language data, or configure the project\'s Node Tesseract backend. The application never installs models during an upload.'),
('Commands everyone should know','UI: python scripts/run_app.py --port 8765. CLI: python scripts/run_project.py data/source/Cadence_Financial_Group_Organizational_Pack.pptx. Tests: python scripts/run_all_tests.py. Acceptance checks: python scripts/run_pipeline_checks.py --runner cadence_detection.runtime:run_for_qa. Use scripts/bootstrap.py to locate optional project-local model dependencies.'),
('Where outputs go','Checked CLI JSON/text goes under output/reports/runs, or a chosen output directory. Raw extraction artifacts remain under data/extracted and are ignored by Git. Web jobs stay in local memory. Do not commit raw evidence, identity mappings, model caches, credentials or extracted personal values. Public release JSON does not contain OCR word text metadata or reversible identity mappings.'),
('Suggested 12-minute demonstration','Person 1: purpose, functional diagram and architecture (2 minutes). Person 2: native text versus OCR and source locations (2). Person 3: categories, context and a negative example (2). Person 4: upload the actual artifacts, inspect warnings and show allowed/blocked downloads (4). Person 5: tests, metrics, limitations and next steps (2). Start the long scanned-policy job before the live walkthrough or keep saved real-run reports as a fallback.'),
('Troubleshooting','No model: install/configure the English NER model; do not bypass the gate. No OCR: configure the local engine/language asset. Port busy: use --port 8766. Exit code 2 from run_project means some input needs review, not necessarily a software crash. Review warnings identify the source location. Never lower thresholds solely to turn a warning into an approval.')])
page('9 | Verification, fixes and project boundaries',[
('Regression evidence',f"Latest regression summary: {tests.get('tests_run','pending')} tests run; {tests.get('passed','pending')} passed; {tests.get('failures','pending')} failures; {tests.get('errors','pending')} errors; {tests.get('skipped','pending')} skipped. The build must be rerun after any remaining failure is fixed. The final JSON report is output/reports/final_tests.json."),
('Integration fixes after the merges','The real runtime adapters now connect all stages. The UI reads actual uploads instead of fixtures. Confidence, document IDs, geometry and metadata are translated. Redaction handles immutable findings, crossing overlaps and invalid spans. Verification checks more than email/phone patterns. Missing dependencies and incomplete extraction block release. Upload authorization, host/origin checks and export enforcement have automated regressions.'),
('What is deliberately blocked','Corrupt, empty, encrypted, spoofed and unsupported files do not release content. Low-confidence/unreadable evidence can be analyzed for local review but cannot be exported as checked text. An image with no readable text does not become proof that no PII is present. Photographs/signatures and original document layers are excluded from the text-only output.'),
('Claims that need independent source review','The case-study targets are not automatically satisfied by merging code or passing unit tests. The complete 80% structure and near-100% real-source PII coverage claims require adjudicated source labels and manual comparison. Report the actual measured scope and denominators. A future unfamiliar format/language or damaged scan can require more recognizers or review.'),
('Responsible deployment','This is a local demonstration and learning application. Internet deployment would require stronger authentication, TLS, storage controls, user isolation, operational monitoring and a retention policy. The local server is intentionally loopback-only. There is no external LLM endpoint in this project; the QA runner uses a local downstream spy.'),
('What each person should be able to answer','What input does your component receive? What output and status does it produce? Which errors does it catch? How is it tested? What could still go wrong? Where is the evidence for your claim? Everyone should be able to explain why a quarantined document must not be presented as safe.')])

doc=Document();sec=doc.sections[0];sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.75);sec.right_margin=Inches(.75)
normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10.5);normal.paragraph_format.space_after=Pt(7)
for h in ['Title','Heading 1','Heading 2']:doc.styles[h].font.color.rgb=RGBColor.from_string('173456')
for i,(title,sections) in enumerate(pages):
    if i:doc.add_page_break()
    doc.add_heading(title,0 if i==0 else 1)
    for heading,text in sections:
        doc.add_heading(heading,2);doc.add_paragraph(text)
footer=sec.footer.paragraphs[0];footer.alignment=2
footer.add_run('Cadence team guide | ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
for node in list(doc.styles.element.xpath('.//w:pBdr')):
    node.getparent().remove(node)
out=ROOT/'output/documents';out.mkdir(parents=True,exist_ok=True)
doc.save(out/'Cadence_Project_and_Team_Guide_Final.docx')

# Editable 16:9 stakeholder slides; diagrams are native shapes, not screenshots.
prs=Presentation();prs.slide_width=PInches(13.333);prs.slide_height=PInches(7.5)
navy=PColor.from_string('15304F');teal=PColor.from_string('087D79');gray=PColor.from_string('43566E')
def text(slide,x,y,w,h,value,size=22,color=gray,bold=False):
    box=slide.shapes.add_textbox(PInches(x),PInches(y),PInches(w),PInches(h));tf=box.text_frame;tf.word_wrap=True
    for i,line in enumerate(value.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.size=PPt(size);p.font.name='Calibri';p.font.bold=bold;p.font.color.rgb=color
    return box
slides=[
('Protecting personal data in Cadence evidence','Local extraction, detection, redaction and verification\nCase Study 2 | 10-15-minute stakeholder walkthrough\n14 October 2026'),
('The requirement and supplied evidence','9 organizational slides: directory, charts and RACI contacts\n35 scanned policy pages: prose, tables and screenshots\nTraining DOCX: 16 tables and 41 media entries\nTargets: preserve ≥80% structure/context; approach complete PII coverage'),
('Functional solution blueprint','diagram'),
('Technical architecture and local boundary','architecture'),
('Extraction logic and source evidence','Use native text where available; OCR page images and embedded media\nPreserve ordered blocks, table indices, source paths and word boxes\nConvert supported SVG and Windows EMF locally\nIncomplete OCR or unsupported content becomes a review warning'),
('PII detection and classification','14 text categories, including names, identifiers and financial references\nCombine local NER with context, validators and document-local aliases\nCount repeated mentions; preserve exact Unicode offsets\nKeep business dates, control IDs and currency as negative test cases'),
('Redaction and release verification','Validate spans; mask unions of overlaps from right to left\nUse stable typed placeholders within each document\nRescan output and verify known original spans were removed\nBlock export on incomplete extraction, missing model or unresolved PII'),
('Reliable outcomes: evidence and boundaries',f"Regression suite: {tests.get('passed','pending')} passed / {tests.get('tests_run','pending')} run\nSupplied synthetic text fixture: 18/18 exact matches, no false positives\nReal files are processed; OCR review can legitimately block export\nFull-source PII recall and ≥80% structure require adjudicated source review"),
('Live walkthrough: the actual upload flow','Start: python scripts/run_app.py --port 8765\nUpload an actual PPTX, PDF or DOCX and follow stage progress\nInspect counts, locations, masked text and OCR warnings\nShow checked-text export for approved jobs and blocked export for review jobs'),
('Team handoff and final decisions','Person 1: lifecycle and architecture | Person 2: extraction evidence\nPerson 3: recognizers | Person 4: release behavior and UI\nPerson 5: independent scoring and presentation evidence\nText-only output; visual PII and uncertain scans remain a review responsibility')]
for i,(title,body) in enumerate(slides):
    slide=prs.slides.add_slide(prs.slide_layouts[6]);text(slide,.6,.4,12.1,1.0,title,32,navy,True)
    if body=='diagram':
        labels=['Upload','Validate','Extract / OCR','Detect PII','Mask spans','Verify / release']
        for j,label in enumerate(labels):
            x=.6+j*2.12;shape=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,PInches(x),PInches(2.1),PInches(1.85),PInches(1.1));shape.fill.solid();shape.fill.fore_color.rgb=navy;shape.line.color.rgb=navy
            text(slide,x+.1,2.35,1.65,.65,label,18,PColor(255,255,255),True)
            if j<5:text(slide,x+1.87,2.35,.27,.5,'→',21,teal,True)
        text(slide,.8,4.0,11.8,1.4,'Completed checks → sanitized plain-text export\nMissing, uncertain or residual content → quarantine and local review',25,teal,True)
    elif body=='architecture':
        components=[(.8,1.8,'Local upload + session checks\nweb.py'),(4.8,1.8,'File checks + manifest\nvalidator.py / runtime.py'),(8.8,1.8,'PDF / OOXML + local OCR\nPDFium / Tesseract'),(8.8,3.85,'NER + contextual recognizers\nspaCy / comprehensive.py'),(4.8,3.85,'Union spans + stable tokens\nredactor.py'),(.8,3.85,'Rescan + completeness\ngate.py + known spans')]
        for x,y,label in components:
            box=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,PInches(x),PInches(y),PInches(3.5),PInches(1.25));box.fill.solid();box.fill.fore_color.rgb=navy;box.line.color.rgb=navy
            text(slide,x+.12,y+.18,3.25,.95,label,17,PColor(255,255,255),True)
        for x in [4.35,8.35]:text(slide,x,2.15,.4,.5,'→',24,teal,True)
        for x in [4.35,8.35]:text(slide,x,4.2,.4,.5,'←',24,teal,True)
        text(slide,12.4,3.15,.4,.5,'↓',24,teal,True)
        text(slide,8.8,3.15,3.5,.6,'Normalized blocks + source locations',13,gray)
        text(slide,.9,5.7,11.5,.6,'Pass → checked plain-text export | otherwise → quarantine + local review',20,teal,True)
        text(slide,.9,6.35,11.5,.4,'Document-local workers and mappings; public audit reports contain counts and locations.',16,gray)
    else:text(slide,.8,1.75,11.7,4.6,body,25)
    text(slide,.8,6.95,11.5,.3,f'Cadence Case Study 2 | {i+1}',11,gray)
    if slide.has_notes_slide:slide.notes_slide.notes_text_frame.text='Explain the component with one concrete example. State measured scope and do not claim universal 100% PII accuracy.'
presentation=ROOT/'presentation/Cadence_Case_Study_Solution.pptx';prs.save(presentation)
print('Created beginner guide:',out/'Cadence_Project_and_Team_Guide_Final.docx')
print('Created stakeholder presentation:',presentation)
