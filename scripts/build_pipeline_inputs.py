"""Generate local synthetic supported, empty, corrupt and spoofed documents."""

import json
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]


def main():
    # Optional fixture-building dependencies only, not evaluator dependencies.
    from docx import Document
    from pptx import Presentation
    from pptx.util import Inches
    from reportlab.pdfgen import canvas
    from pypdf import PdfReader, PdfWriter
    out = ROOT / "data/synthetic/pipeline_inputs"
    out.mkdir(parents=True, exist_ok=True)
    text = "Privacy review. Noor Example, noor@example.invalid. Control GRP-POL-001 remains applicable."
    d = Document()
    d.add_heading("Privacy review", level=1)
    d.add_paragraph(text)
    d.save(out / "valid.docx")
    p = Presentation()
    slide = p.slides.add_slide(p.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2)).text_frame.text = text
    p.save(out / "valid.pptx")
    c = canvas.Canvas(str(out / "valid.pdf"))
    c.drawString(40, 780, text)
    c.save()
    for ext in ("pdf", "pptx", "docx"):
        (out / ("empty." + ext)).write_bytes(b"")
        (out / ("corrupt." + ext)).write_bytes(b"invalid document container")
    (out / "unsupported.txt").write_text(text, encoding="utf-8")
    (out / "spoofed.pdf").write_bytes((out / "valid.docx").read_bytes())
    writer = PdfWriter()
    for page in PdfReader(out / "valid.pdf").pages:
        writer.add_page(page)
    writer.encrypt("synthetic-test-password")
    with (out / "encrypted.pdf").open("wb") as f:
        writer.write(f)
    print("Generated 12 local pipeline inputs. No real personal data.")


if __name__ == "__main__":
    main()
