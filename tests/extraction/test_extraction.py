"""Deterministic synthetic tests; real OCR is separately exercised on supplied sources."""
from __future__ import annotations

import contextlib
import io
import json
import os
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile

sys.path.insert(0, str(Path(__file__).resolve().parents[2]/"src"))
from cadence_detection.extraction import extract_document, ExtractionOptions
from cadence_detection.extraction.models import ExtractionResult
from cadence_detection.extraction.ocr import OCRLine, OCRError, TesseractOCR, NodeTesseractOCR, parse_tsv
from cadence_detection.extraction.images import recognize_image, open_image
from cadence_detection.extraction.office import paragraph_text, xml, relationships, NS
from cadence_detection.extraction.cli import main
from PIL import Image, ImageDraw
from lxml import etree
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.util import Inches
from reportlab.pdfgen.canvas import Canvas
from reportlab.lib.utils import ImageReader
from pypdf import PdfReader, PdfWriter


class FakeOCR:
    def __init__(self, text="Image Contact", confidence=96, fail=False, empty=False):
        self.text, self.confidence, self.fail, self.empty = text, confidence, fail, empty
        self.calls = []

    def recognize(self, image, *, sparse=False):
        self.calls.append((image.size, sparse))
        if self.fail:
            raise OCRError("Injected failure with private text")
        if self.empty:
            return []
        return [OCRLine(self.text, [0, 0, image.width/2, image.height/2], self.confidence)]

    def close(self):
        pass


class ExtractionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.ocr = FakeOCR()
        self.image = Image.new("RGB", (400, 120), "white")
        ImageDraw.Draw(self.image).text((10, 10), "CONTACT", fill="black")

    def tearDown(self):
        self.temp.cleanup()

    def extract(self, path, **kwargs):
        return extract_document(path, ocr=self.ocr, **kwargs)

    def docx(self):
        d = Document(); d.add_paragraph("First paragraph")
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text = "Name"; t.cell(0, 1).text = "Email"
        t.cell(1, 0).text = "Anne O'Neil"; t.cell(1, 1).text = "anne@example.test"
        d.add_paragraph("After table")
        path = self.root/"sample.docx"; d.save(path)
        return path

    def pptx(self):
        p = Presentation(); s = p.slides.add_slide(p.slide_layouts[6])
        box = s.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
        box.text = "Slide One"
        s.notes_slide.notes_text_frame.text = "Private note contact"
        s2 = p.slides.add_slide(p.slide_layouts[6])
        s2.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text = "Slide Two"
        path = self.root/"sample.pptx"; p.save(path)
        return path

    def pdf(self, *, image=False, mixed=False, blank=False, pages=1):
        path = self.root/"sample.pdf"; c = Canvas(str(path), pagesize=(400, 300))
        for n in range(pages):
            if not blank and (not image or mixed):
                c.drawString(25, 270, f"Native page {n+1}")
                c.drawString(25, 245, "Context next line")
            if image:
                c.drawImage(ImageReader(self.image), 25, 100, width=300, height=90)
            c.showPage()
        c.save(); return path

    def rewrite(self, path, modify):
        with zipfile.ZipFile(path) as z:
            data = {n: z.read(n) for n in z.namelist()}
        modify(data)
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
            for name, value in data.items():
                z.writestr(name, value)

    def codes(self, result):
        return {i.code for i in result.issues}

    def test_missing_file(self):
        r = self.extract(self.root/"absent.pdf"); self.assertEqual(r.processing_status, "failed")
        self.assertIn("FILE_NOT_FOUND", self.codes(r))

    def test_directory(self):
        self.assertIn("FILE_NOT_FOUND", self.codes(self.extract(self.root)))

    def test_empty_file(self):
        p = self.root/"empty.pdf"; p.touch()
        self.assertIn("FILE_LIMIT", self.codes(self.extract(p)))

    def test_unsupported_extension(self):
        p = self.root/"test.txt"; p.write_text("data")
        self.assertIn("UNSUPPORTED_FORMAT", self.codes(self.extract(p)))

    def test_wrong_pdf_signature(self):
        p = self.root/"test.pdf"; p.write_text("not a pdf")
        self.assertIn("PARSE_FAILED", self.codes(self.extract(p)))

    def test_corrupt_office_package(self):
        p = self.root/"test.docx"; p.write_bytes(b"PK invalid")
        self.assertEqual(self.extract(p).processing_status, "failed")

    def test_mismatched_office_extension(self):
        p = self.docx(); renamed = p.with_suffix(".pptx"); p.rename(renamed)
        self.assertIn("PARSE_FAILED", self.codes(self.extract(renamed)))

    def test_uppercase_extension(self):
        p = self.docx(); q = p.with_suffix(".DOCX"); p.rename(q)
        self.assertTrue(self.extract(q).extraction_complete)

    def test_file_size_limit(self):
        self.assertIn("FILE_LIMIT", self.codes(self.extract(self.docx(), options=ExtractionOptions(max_file_bytes=10))))

    def test_unreadable_input(self):
        p = self.docx()
        with patch("pathlib.Path.open", side_effect=PermissionError):
            r = self.extract(p)
        self.assertIn("FILE_UNREADABLE", self.codes(r))

    def test_zip_expansion_limit(self):
        self.assertIn("PACKAGE_LIMIT", self.codes(self.extract(self.docx(), options=ExtractionOptions(max_uncompressed_bytes=10))))

    def test_zip_entry_limit(self):
        self.assertIn("PACKAGE_LIMIT", self.codes(self.extract(self.docx(), options=ExtractionOptions(max_zip_entries=1))))

    def test_duplicate_zip_members(self):
        p = self.docx()
        import warnings
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            with zipfile.ZipFile(p, "a") as z:
                z.writestr("word/document.xml", "duplicate")
        self.assertIn("PARSE_FAILED", self.codes(self.extract(p)))

    def test_xxe_xml_rejected(self):
        p = self.docx()
        self.rewrite(p, lambda data: data.update({"word/document.xml": b'<!DOCTYPE x [<!ENTITY y SYSTEM "file:///secret">]><x>&y;</x>'}))
        self.assertIn("PARSE_FAILED", self.codes(self.extract(p)))

    def test_invalid_options(self):
        for options in ({"dpi": 10}, {"dpi": 900}, {"min_ocr_confidence": -1}, {"max_units": 0}, {"max_image_pixels": 0}):
            with self.subTest(options=options), self.assertRaises(ValueError):
                ExtractionOptions(**options)

    def test_docx_paragraph_table_order(self):
        r = self.extract(self.docx())
        body = [b.text for b in r.blocks if b.source_kind in ("paragraph", "table_cell")]
        self.assertEqual(body, ["First paragraph", "Name", "Email", "Anne O'Neil", "anne@example.test", "After table"])
        self.assertTrue(r.extraction_complete)

    def test_table_coordinates(self):
        r = self.extract(self.docx())
        cell = next(b for b in r.blocks if b.text == "anne@example.test")
        self.assertEqual((cell.table_row, cell.table_column), (1, 1))
        self.assertIn(":table:1:row:1:column:1:", cell.source_location)

    def test_docx_no_invented_page_number(self):
        r = self.extract(self.docx())
        self.assertTrue(all(b.page_or_slide is None for b in r.blocks))
        self.assertIn("WORD_PAGINATION", self.codes(r))

    def test_unicode_tabs_and_breaks_preserved(self):
        p = self.docx(); d = Document(p)
        d.add_paragraph("Élodie 李\tcontext\nnext"); d.save(p)
        self.assertIn("Élodie 李\tcontext\nnext", [b.text for b in self.extract(p).blocks])

    def test_header_footer(self):
        p = self.docx(); d = Document(p)
        d.sections[0].header.paragraphs[0].text = "Header contact"
        d.sections[0].footer.paragraphs[0].text = "Footer contact"; d.save(p)
        r = self.extract(p)
        self.assertTrue(any(b.text == "Header contact" and b.source_kind == "header" for b in r.blocks))
        self.assertTrue(any(b.text == "Footer contact" and b.source_kind == "footer" for b in r.blocks))

    def test_nested_table(self):
        p = self.docx(); d = Document(p)
        t = d.tables[0].cell(1, 0).add_table(rows=1, cols=1); t.cell(0, 0).text = "Nested contact"; d.save(p)
        r = self.extract(p); b = next(b for b in r.blocks if b.text == "Nested contact")
        self.assertEqual(b.source_location.count(":table:"), 2)

    def test_horizontal_merge_grid_column(self):
        p = self.docx(); d = Document(p)
        d.tables[0].cell(0, 0).merge(d.tables[0].cell(0, 1)); d.save(p)
        r = self.extract(p)
        self.assertEqual(sum(b.text == "Name" for b in r.blocks), 1)

    def test_textbox_not_duplicated_or_dropped(self):
        p = self.docx()
        def modify(data):
            root = etree.fromstring(data["word/document.xml"])
            para = root.find("w:body/w:p", NS)
            content = etree.SubElement(etree.SubElement(para, qn("w:pict")), qn("w:txbxContent"))
            t = etree.SubElement(etree.SubElement(etree.SubElement(content, qn("w:p")), qn("w:r")), qn("w:t")); t.text = "Textbox Contact"
            data["word/document.xml"] = etree.tostring(root)
        self.rewrite(p, modify); r = self.extract(p)
        self.assertEqual(sum(b.text == "Textbox Contact" for b in r.blocks), 1)
        self.assertTrue(any(b.text == "Textbox Contact" and b.source_kind == "textbox" for b in r.blocks))

    def test_footnotes_comments_endnotes(self):
        for part, container in (("footnotes", "footnote"), ("comments", "comment"), ("endnotes", "endnote")):
            with self.subTest(part=part):
                p = self.docx()
                content = f'<w:{part} xmlns:w="{NS["w"]}"><w:{container}><w:p><w:r><w:t>Extra contact</w:t></w:r></w:p></w:{container}></w:{part}>'
                self.rewrite(p, lambda data: data.update({f"word/{part}.xml": content.encode()}))
                r = self.extract(p)
                self.assertTrue(any(b.text == "Extra contact" and b.source_kind == part for b in r.blocks))

    def test_hidden_and_tracked_content_flagged(self):
        p = self.docx(); d = Document(p)
        d.add_paragraph().add_run("Hidden Contact").font.hidden = True; d.save(p)
        self.assertIn("OFFICE_REVISIONS_HIDDEN", self.codes(self.extract(p)))

    def test_field_instructions_recovered(self):
        p = self.docx()
        def modify(data):
            root = etree.fromstring(data["word/document.xml"])
            para = root.find("w:body/w:p", NS)
            etree.SubElement(etree.SubElement(para, qn("w:r")), qn("w:instrText")).text = 'HYPERLINK "mailto:field@example.test"'
            data["word/document.xml"] = etree.tostring(root)
        self.rewrite(p, modify)
        self.assertTrue(any(b.source_kind == "field_instruction" and "field@example.test" in b.text for b in self.extract(p).blocks))

    def test_tracked_deletion_requires_review(self):
        p = self.docx()
        def modify(data):
            root = etree.fromstring(data["word/document.xml"])
            para = root.find("w:body/w:p", NS)
            deleted = etree.SubElement(para, qn("w:del"))
            etree.SubElement(etree.SubElement(deleted, qn("w:r")), qn("w:delText")).text = "Deleted Contact"
            data["word/document.xml"] = etree.tostring(root)
        self.rewrite(p, modify)
        self.assertIn("OFFICE_REVISIONS_HIDDEN", self.codes(self.extract(p)))

    def test_internal_relationship_path_traversal_rejected(self):
        p = self.docx()
        def modify(data):
            name = "word/_rels/document.xml.rels"; root = etree.fromstring(data[name])
            etree.SubElement(root, "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
                             Id="escape", Type=NS["r"]+"/image", Target="../../outside.png")
            data[name] = etree.tostring(root)
        self.rewrite(p, modify)
        self.assertIn("PARSE_FAILED", self.codes(self.extract(p)))

    def test_external_hyperlink_value_recovered(self):
        p = self.docx()
        def modify(data):
            name = "word/_rels/document.xml.rels"; root = etree.fromstring(data[name])
            etree.SubElement(root, "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
                             Id="mail", Type=NS["r"]+"/hyperlink", Target="mailto:person@example.test", TargetMode="External")
            data[name] = etree.tostring(root)
        self.rewrite(p, modify)
        self.assertTrue(any(b.text == "mailto:person@example.test" for b in self.extract(p).blocks))

    def test_metadata_extract_and_skip(self):
        p = self.docx(); d = Document(p); d.core_properties.author = "Metadata Contact"; d.save(p)
        self.assertTrue(any(b.text == "Metadata Contact" for b in self.extract(p).blocks))
        r = self.extract(p, options=ExtractionOptions(include_metadata=False))
        self.assertFalse(any(b.source_kind == "metadata" for b in r.blocks))
        self.assertFalse(r.extraction_complete)

    def test_docx_picture_ocr_and_provenance(self):
        p = self.docx(); pic = self.root/"picture.png"; self.image.save(pic)
        d = Document(p); d.add_picture(str(pic)); d.save(p)
        r = self.extract(p); b = next(b for b in r.blocks if b.text == "Image Contact")
        self.assertTrue(b.attributes["media_part"].startswith("word/media/"))
        self.assertEqual(b.coordinate_space, "image_pixels_top_left")

    def test_repeated_picture_occurrences_retained(self):
        p = self.docx(); pic = self.root/"picture.png"; self.image.save(pic)
        d = Document(p); d.add_picture(str(pic)); d.add_picture(str(pic)); d.save(p)
        self.assertEqual(sum(b.text == "Image Contact" for b in self.extract(p).blocks), 2)

    def test_unplaced_unsupported_media_quarantined(self):
        p = self.docx(); self.rewrite(p, lambda data: data.update({"word/media/orphan.emf": b"vector"}))
        r = self.extract(p)
        self.assertIn("OFFICE_UNPLACED_MEDIA", self.codes(r)); self.assertIn("OFFICE_IMAGE_UNSUPPORTED", self.codes(r))

    def test_external_image_not_fetched(self):
        p = self.docx()
        def modify(data):
            name = "word/_rels/document.xml.rels"; root = etree.fromstring(data[name])
            etree.SubElement(root, "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
                             Id="external", Type=NS["r"]+"/image", Target="https://example.invalid/private.png", TargetMode="External")
            data[name] = etree.tostring(root)
        self.rewrite(p, modify)
        self.assertIn("OFFICE_EXTERNAL_RESOURCE", self.codes(self.extract(p)))

    def test_pptx_slide_order_and_notes(self):
        r = self.extract(self.pptx())
        self.assertEqual([b.text for b in r.blocks if b.source_kind == "slide_text"], ["Slide One", "Slide Two"])
        self.assertTrue(any(b.text == "Private note contact" and b.source_kind == "speaker_notes" for b in r.blocks))
        self.assertEqual([u["index"] for u in r.units], [1, 2])

    def test_pptx_notes_optout_quarantines(self):
        r = self.extract(self.pptx(), options=ExtractionOptions(include_notes=False))
        self.assertIn("NOTES_SKIPPED", self.codes(r)); self.assertFalse(r.extraction_complete)

    def test_pptx_table(self):
        p = self.pptx(); d = Presentation(p)
        table = d.slides[0].shapes.add_table(2, 2, Inches(1), Inches(3), Inches(4), Inches(1)).table
        table.cell(1, 1).text = "RACI Contact"; d.save(p)
        b = next(b for b in self.extract(p).blocks if b.text == "RACI Contact")
        self.assertEqual((b.table_row, b.table_column, b.page_or_slide), (1, 1, 1))

    def test_pptx_group(self):
        p = self.pptx(); d = Presentation(p)
        group = d.slides[0].shapes.add_group_shape()
        group.shapes.add_textbox(Inches(2), Inches(2), Inches(1), Inches(1)).text = "Grouped Contact"; d.save(p)
        b = next(b for b in self.extract(p).blocks if b.text == "Grouped Contact")
        self.assertIsNotNone(b.bbox); self.assertEqual(b.coordinate_space, "slide_emu_top_left")
        self.assertIn(":group:shape:", b.source_location)

    def test_pptx_picture_position_mapping(self):
        p = self.pptx(); d = Presentation(p); pic = self.root/"p.png"; self.image.save(pic)
        d.slides[0].shapes.add_picture(str(pic), Inches(1), Inches(2), Inches(4), Inches(1.2)); d.save(p)
        b = next(b for b in self.extract(p).blocks if b.text == "Image Contact")
        self.assertEqual(b.bbox[:2], [float(Inches(1)), float(Inches(2))])
        self.assertAlmostEqual(b.bbox[2], float(Inches(3)))

    def test_pptx_hidden_slide_included(self):
        p = self.pptx(); d = Presentation(p); d.slides[0]._element.set("show", "0"); d.save(p)
        r = self.extract(p)
        self.assertIn("PPT_HIDDEN_SLIDE", self.codes(r)); self.assertIn("Slide One", [b.text for b in r.blocks])

    def test_pptx_unit_limit(self):
        self.assertIn("UNIT_LIMIT", self.codes(self.extract(self.pptx(), options=ExtractionOptions(max_units=1))))

    def test_rotated_group_flagged(self):
        p = self.pptx(); d = Presentation(p); group = d.slides[0].shapes.add_group_shape()
        group.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1)).text = "Rotated contact"
        group.rotation = 45; d.save(p)
        self.assertIn("PPT_GROUP_TRANSFORM", self.codes(self.extract(p)))

    def test_missing_slide_relationship_does_not_drop_other_slide(self):
        p = self.pptx()
        def modify(data):
            root = etree.fromstring(data["ppt/_rels/presentation.xml.rels"])
            for rel in list(root):
                if rel.get("Target") == "slides/slide1.xml":
                    root.remove(rel)
            data["ppt/_rels/presentation.xml.rels"] = etree.tostring(root)
        self.rewrite(p, modify); r = self.extract(p)
        self.assertIn("PPT_SLIDE_FAILED", self.codes(r))
        self.assertTrue(any(b.text == "Slide Two" for b in r.blocks))
        self.assertFalse(r.units[0]["processed"]); self.assertTrue(r.units[1]["processed"])

    def test_native_pdf_and_boxes(self):
        r = self.extract(self.pdf()); self.assertTrue(r.extraction_complete)
        self.assertTrue(any("Native page 1" in b.text for b in r.blocks))
        self.assertFalse(self.ocr.calls)
        self.assertTrue(all(b.bbox for b in r.blocks if b.source_kind == "pdf_text"))

    def test_scanned_pdf_uses_ocr(self):
        r = self.extract(self.pdf(image=True))
        self.assertEqual(r.units[0]["method"], "page_ocr")
        self.assertTrue(any(b.text == "Image Contact" for b in r.blocks))

    def test_mixed_pdf_keeps_native_and_image_text(self):
        r = self.extract(self.pdf(image=True, mixed=True))
        self.assertTrue(any("Native page" in b.text for b in r.blocks))
        self.assertTrue(any(b.text == "Image Contact" for b in r.blocks))

    def test_pdf_page_order(self):
        r = self.extract(self.pdf(pages=3)); self.assertEqual([u["index"] for u in r.units], [1, 2, 3])
        self.assertTrue(any("Native page 3" in b.text and b.page_or_slide == 3 for b in r.blocks))

    def test_pdf_page_limit(self):
        r = self.extract(self.pdf(pages=2), options=ExtractionOptions(max_units=1))
        self.assertIn("UNIT_LIMIT", self.codes(r)); self.assertFalse(r.extraction_complete)

    def test_ocr_detail_crop_and_source_mapping(self):
        r = self.extract(self.pdf(image=True), options=ExtractionOptions(ocr_regions={1: [[0.25, 0.25, 0.75, 0.75]]}))
        b = next(b for b in r.blocks if b.attributes.get("supplementary_region"))
        self.assertAlmostEqual(b.bbox[0], 100, delta=1)
        self.assertAlmostEqual(b.bbox[1], 75, delta=1)
        self.assertEqual(r.units[0]["detail_region_blocks"], 1)

    def test_invalid_crop_configuration(self):
        for regions in ({0: [[0, 0, 1, 1]]}, {1: [[0, 0, 2, 1]]}, {1: [[1, 0, 0, 1]]}):
            with self.subTest(regions=regions), self.assertRaises(ValueError):
                ExtractionOptions(ocr_regions=regions)

    def test_nonexistent_crop_page_flagged(self):
        r = self.extract(self.pdf(), options=ExtractionOptions(ocr_regions={99: [[0, 0, 1, 1]]}))
        self.assertIn("OCR_REGION_PAGE", self.codes(r))

    def test_pdf_pixel_limit(self):
        r = self.extract(self.pdf(image=True), options=ExtractionOptions(max_image_pixels=100))
        self.assertIn("IMAGE_LIMIT", self.codes(r)); self.assertFalse(r.units[0]["processed"])

    def test_pdf_encrypted_password(self):
        p = self.pdf(); reader = PdfReader(p); w = PdfWriter(); w.append(reader); w.encrypt("secret")
        target = self.root/"encrypted.pdf"
        with target.open("wb") as stream:
            w.write(stream)
        for password in (None, "wrong"):
            r = self.extract(target, options=ExtractionOptions(password=password))
            self.assertIn("PDF_PASSWORD", self.codes(r))
        self.assertTrue(self.extract(target, options=ExtractionOptions(password="secret")).extraction_complete)

    def test_pdf_annotation_flagged(self):
        p = self.root/"annotated.pdf"; c = Canvas(str(p)); c.drawString(50, 700, "Visible")
        c.textAnnotation("Hidden contact", Rect=(50, 500, 100, 550)); c.save()
        self.assertIn("PDF_ANNOTATIONS", self.codes(self.extract(p)))

    def test_pdf_glyph_rectangles_preserve_contact_context(self):
        p = self.root/"glyphs.pdf"
        c = Canvas(str(p))
        lines = ["Full Name: Alex Example", "Phone Number: +91 9876543210", "Email Address: alex@example.test"]
        for row, line in enumerate(lines):
            x = 40
            for character in line:
                c.drawString(x, 700 - row * 20, character)
                x += c.stringWidth(character, "Helvetica", 12)
        c.save()
        result = self.extract(p)
        blocks = [b for b in result.blocks if b.source_kind == "pdf_text"]
        self.assertEqual(len(blocks), 1)
        for line in lines:
            self.assertIn(line, blocks[0].text)
        self.assertTrue(result.extraction_complete)

    def test_pdf_uri_link_target_extracted(self):
        p = self.root/"link.pdf"
        c = Canvas(str(p)); c.drawString(50, 700, "Contact")
        c.linkURL("mailto:alex@example.test", (50, 690, 100, 710)); c.save()
        result = self.extract(p)
        self.assertTrue(result.extraction_complete)
        self.assertNotIn("PDF_ANNOTATIONS", self.codes(result))
        self.assertIn("mailto:alex@example.test", [b.text for b in result.blocks if b.source_kind == "pdf_link"])

    def test_pdf_attachment_flagged(self):
        reader = PdfReader(self.pdf()); w = PdfWriter(); w.append(reader); w.add_attachment("secret.txt", b"Contact")
        p = self.root/"attached.pdf"
        with p.open("wb") as stream:
            w.write(stream)
        self.assertIn("PDF_ATTACHMENTS", self.codes(self.extract(p)))

    def test_native_pdf_metadata_optout_flagged(self):
        self.assertIn("METADATA_SKIPPED", self.codes(self.extract(self.pdf(), options=ExtractionOptions(include_metadata=False))))

    def test_blank_pdf_not_safe_success(self):
        self.ocr = FakeOCR(empty=True); r = self.extract(self.pdf(blank=True))
        self.assertFalse(r.extraction_complete)

    def test_rotated_scan(self):
        reader = PdfReader(self.pdf(image=True)); w = PdfWriter(); page = reader.pages[0]; page.rotate(90); w.add_page(page)
        p = self.root/"rotated.pdf"
        with p.open("wb") as stream:
            w.write(stream)
        r = self.extract(p); b = next(b for b in r.blocks if b.source_kind == "ocr")
        self.assertEqual(b.attributes["page_rotation"], 90)
        self.assertEqual(b.coordinate_space, "pdf_points_top_left_displayed")

    def test_ocr_failure_retains_native_text(self):
        self.ocr = FakeOCR(fail=True); r = self.extract(self.pdf(image=True, mixed=True))
        self.assertIn("OCR_FAILED", self.codes(r)); self.assertFalse(r.extraction_complete)
        self.assertTrue(any("Native page" in b.text for b in r.blocks))
        self.assertNotIn("private text", json.dumps(r.to_dict()))

    def test_low_ocr_confidence_quarantines_all_blocks(self):
        self.ocr = FakeOCR(confidence=20); r = self.extract(self.pdf(image=True, mixed=True))
        self.assertIn("LOW_OCR_CONFIDENCE", self.codes(r))
        self.assertTrue(all(not b.extraction_complete for b in r.blocks))

    def test_no_ocr_engine(self):
        with patch("cadence_detection.extraction.pipeline.default_ocr", return_value=None):
            r = extract_document(self.pdf(image=True))
        self.assertIn("OCR_UNAVAILABLE", self.codes(r))

    def test_empty_ocr_on_nonblank_image(self):
        self.ocr = FakeOCR(empty=True); r = self.extract(self.pdf(image=True))
        self.assertIn("IMAGE_NO_TEXT", self.codes(r))

    def test_json_schema_stable_ids_and_order(self):
        p = self.docx(); a, b = self.extract(p), self.extract(p)
        self.assertEqual(a.to_dict(), b.to_dict())
        self.assertEqual([x.reading_order for x in a.blocks], list(range(len(a.blocks))))
        self.assertEqual(len({x.block_id for x in a.blocks}), len(a.blocks))
        self.assertIn("coverage", json.loads(json.dumps(a.to_dict())))

    def test_cli_exit_code_and_raw_json(self):
        p = self.docx(); output = self.root/"out"; capture = io.StringIO()
        with contextlib.redirect_stdout(capture), patch("cadence_detection.extraction.cli.default_ocr", return_value=self.ocr):
            code = main([str(p), "--output-dir", str(output)])
        self.assertEqual(code, 0); self.assertNotIn("Anne", capture.getvalue())
        self.assertEqual(len(list(output.glob("*.json"))), 1)
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(main([str(self.root/"missing.pdf"), "--output-dir", str(output)]), 2)


class ImageAndOCRTests(unittest.TestCase):
    def result(self):
        return ExtractionResult("d", "test", "png", "hash")

    def test_tsv_hierarchy_and_confidence(self):
        text = "level\tpage_num\tblock_num\tpar_num\tline_num\tword_num\tleft\ttop\twidth\theight\tconf\ttext\n"
        text += "5\t1\t1\t1\t1\t1\t10\t20\t40\t10\t90\tAnne\n5\t1\t1\t1\t1\t2\t60\t20\t20\t10\t100\tLi\n"
        rows = parse_tsv(text); self.assertEqual(rows[0].text, "Anne Li")
        self.assertEqual(rows[0].bbox, [10, 20, 80, 30]); self.assertAlmostEqual(rows[0].confidence, 93.3333333)
        self.assertEqual(rows[0].attributes["words"][1]["start"], 5)
        self.assertEqual(rows[0].text[5:7], "Li")

    def test_bad_tsv_rows_ignored(self):
        text = "level\tconf\ttext\n5\tinvalid\tprivate\n4\t99\tline\n5\t-1\tignored\n"
        self.assertEqual(parse_tsv(text), [])

    def test_nonfinite_tsv_coordinates_rejected(self):
        self.assertEqual(parse_tsv("5\t1\t1\t1\t1\t1\tnan\t20\t40\t10\t91\tinvalid\n"), [])

    def test_tesseract_js_headerless_tsv(self):
        text = "5\t1\t1\t1\t1\t1\t10\t20\t40\t10\t91.5\tSynthetic\n"
        rows = parse_tsv(text)
        self.assertEqual(rows[0].text, "Synthetic")
        self.assertEqual(rows[0].confidence, 91.5)

    def test_svg_external_reference_rejected(self):
        svg = b'<svg xmlns="http://www.w3.org/2000/svg"><image href="https://example.invalid/x.png"/></svg>'
        with self.assertRaises(ValueError):
            open_image(svg, ".svg", 100000)

    def test_svg_doctype_rejected(self):
        with self.assertRaises(ValueError):
            open_image(b'<!DOCTYPE svg><svg/>', ".svg", 100000)

    def test_svg_css_import_rejected(self):
        with self.assertRaises(ValueError):
            open_image(b'<svg xmlns="http://www.w3.org/2000/svg"><style>@import "https://example.invalid";</style></svg>', ".svg", 100000)

    def test_image_upscale_boxes_map_back(self):
        result = self.result(); image = Image.new("RGB", (100, 50), "white"); ocr = FakeOCR()
        recognize_image(image, result=result, options=ExtractionOptions(), ocr=ocr, location="img")
        self.assertEqual(ocr.calls[0][0], (300, 150)); self.assertEqual(result.blocks[0].bbox, [0, 0, 50, 25])

    def test_image_missing_engine(self):
        r = self.result(); recognize_image(Image.new("RGB", (10, 10)), result=r, options=ExtractionOptions(), ocr=None, location="image")
        self.assertEqual(r.issues[0].code, "OCR_UNAVAILABLE")

    def test_corrupt_emf_is_explicitly_rejected(self):
        with self.assertRaises(Exception):
            open_image(b"emf", ".emf", 1000)

    def test_corrupt_image(self):
        with self.assertRaises(Exception):
            open_image(b"invalid", ".png", 1000)

    def test_image_pixel_limit(self):
        b = io.BytesIO(); Image.new("RGB", (100, 100)).save(b, format="PNG")
        with self.assertRaises(ValueError):
            open_image(b.getvalue(), ".png", 100)

    def test_alpha_composited_on_white(self):
        b = io.BytesIO(); Image.new("RGBA", (10, 10), (0, 0, 0, 0)).save(b, format="PNG")
        self.assertEqual(open_image(b.getvalue(), ".png", 1000).getpixel((0, 0)), (255, 255, 255))

    def test_multiframe_is_explicitly_flagged(self):
        b = io.BytesIO(); a = Image.new("RGB", (10, 10), "white"); second = Image.new("RGB", (10, 10), "black")
        a.save(b, format="GIF", save_all=True, append_images=[second])
        with self.assertRaises(ValueError):
            open_image(b.getvalue(), ".gif", 1000)

    def test_tesseract_missing_engine(self):
        with self.assertRaises(OCRError):
            TesseractOCR(executable="nonexistent-cadence-tesseract").recognize(Image.new("RGB", (5, 5)))

    def test_node_missing_model(self):
        with tempfile.TemporaryDirectory() as directory:
            with self.assertRaises(OCRError):
                NodeTesseractOCR("node", "tesseract.js", directory).recognize(Image.new("RGB", (5, 5)))


if __name__ == "__main__":
    unittest.main()
